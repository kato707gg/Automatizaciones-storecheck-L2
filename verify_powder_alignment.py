import datetime
import shutil
import sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

sys.path.insert(0, r"c:\Users\katoH\OneDrive\Documentos\Proyectos personales\Automatizaciones storecheck L2")

from core.catalogacion.completa import procesar_matriz

source = r"C:\Users\katoH\Downloads\MATRIZ DE CATALOGACÓN actual 1 (4) (1) - copia.xlsx"
products = r"C:\Users\katoH\Downloads\layout_products_188865 (5).xlsx"
places = r"C:\Users\katoH\Downloads\layout_places_188865 (5).xlsx"
base = r"c:\Users\katoH\OneDrive\Documentos\Proyectos personales\Automatizaciones storecheck L2"

ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
target = rf"C:\Users\katoH\Downloads\Matriz_FIX2_{ts}.xlsx"
shutil.copy(source, target)
print("FILE", target)

ok = procesar_matriz(target, products, places, base)
print("OK", ok)

wb = load_workbook(target, data_only=False)
ws_prod = wb["PRODUCTOS"]
ws_cfg = wb["CONFIGURACIÓN DE ANAQUEL"]
ws_orig = load_workbook(source, data_only=False)["CONFIGURACIÓN DE ANAQUEL"]

print("PRODUCTOS BY-CF")
for c in range(column_index_from_string("BY"), column_index_from_string("CF") + 1):
    print(get_column_letter(c), [ws_prod.cell(row=r, column=c).value for r in (4, 5, 6)])

headers = [
    "ELECTROLIFE ZERO POLVO UVA 12 SOBRES",
    "ELECTROLIFE ZERO POLVO NARANJA 12 SOBRES",
    "ELECTROLIFE ZERO POLVO FRESA KIWI 12 SOBRES",
]

def find_col(ws, rows, text):
    for row in rows:
        for col in range(1, ws.max_column + 1):
            v = ws.cell(row=row, column=col).value
            if v and str(v).strip().upper() == text:
                return col
    return None

for text in headers:
    col_orig = find_col(ws_orig, (3, 4), text)
    col_proc = find_col(ws_cfg, (3,), text)
    ones_orig = sum(1 for r in range(5, ws_orig.max_row + 1) if col_orig and ws_orig.cell(row=r, column=col_orig).value == 1)
    ones_proc = sum(1 for r in range(5, ws_cfg.max_row + 1) if col_proc and ws_cfg.cell(row=r, column=col_proc).value == 1)
    print("COUNTS", text, get_column_letter(col_orig) if col_orig else None, ones_orig, get_column_letter(col_proc) if col_proc else None, ones_proc)
