from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

path = r'C:\Users\katoH\Downloads\Matriz_Aqualit_CF_20260610_044753.xlsx'
wb = load_workbook(path, data_only=False)
needle = 'ELECTROLIFE ZERO POLVO NARANJA'

for sheet_name, row, start in [
    ('CONFIGURACIÓN DE ANAQUEL', 3, column_index_from_string('K')),
    ('PRODUCTOS', 4, 1),
]:
    print(sheet_name)
    ws = wb[sheet_name]
    found = False
    for col in range(start, min(ws.max_column + 1, 260)):
        value = ws.cell(row=row, column=col).value
        if value and needle in str(value).upper():
            print(get_column_letter(col), col, value, ws.cell(row=2, column=col).value)
            found = True
            break
    if not found:
        print('NOT FOUND')
