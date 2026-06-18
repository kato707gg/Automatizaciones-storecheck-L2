# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, r"c:\Users\katoH\OneDrive\Documentos\Proyectos personales\Automatizaciones storecheck L2")

# Limpiar módulos
for mod in list(sys.modules.keys()):
    if any(x in mod for x in ['catalogacion', 'completa', 'producto_general']):
        del sys.modules[mod]

from core.catalogacion.completa import procesar_matriz
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

print("="*100)
print("PROCESAMIENTO CON ANÁLISIS DE POLVO NARANJA")
print("="*100)

# Primero, analizar el original
print("\n[1] Analizando MATRIZ ORIGINAL...")
wb_orig = load_workbook(r"C:\Users\katoH\Downloads\MATRIZ DE CATALOGACÓN actual 1 (4) (1) - copia.xlsx")
ws_orig = wb_orig["CONFIGURACIÓN DE ANAQUEL"]

print(f"  Buscando ELECTROLIFE ZERO POLVO NARANJA en fila 4...")
for col in range(1, ws_orig.max_column + 1):
    val = ws_orig.cell(row=4, column=col).value
    if val and "NARANJA" in str(val).upper() and "POLVO" in str(val).upper():
        col_letter = get_column_letter(col)
        print(f"  ✓ ENCONTRADO en {col_letter} ({col}): {val}")
        print(f"    Fila 3 (SKU): {ws_orig.cell(row=3, column=col).value}")
        print(f"    Fila 5 (dato): {ws_orig.cell(row=5, column=col).value}")
        break

# Ahora procesar
print("\n[2] Ejecutando pipeline...")
ruta_archivo = r"C:\Users\katoH\Downloads\MATRIZ DE CATALOGACÓN actual 1 (4) (1) - copia.xlsx"
ruta_products = r"C:\Users\katoH\Downloads\layout_products_188865 (5).xlsx"
ruta_places = r"C:\Users\katoH\Downloads\layout_places_188865 (5).xlsx"
ruta_base = r"C:\Users\katoH\OneDrive\Documentos\Proyectos personales\Automatizaciones storecheck L2"

# Crear archivo nuevo basado en la copia para no modificar el respaldo
import shutil
import datetime

timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
ruta_nueva = f"C:/Users/katoH/Downloads/Matriz_Analisis_{timestamp}.xlsx"
print(f"  Creando copia de trabajo: {ruta_nueva}")
shutil.copy(ruta_archivo, ruta_nueva)

result = procesar_matriz(ruta_nueva, ruta_products, ruta_places, ruta_base)

# Analizar el resultado
print("\n[3] Analizando MATRIZ PROCESADA...")
try:
    wb_proc = load_workbook(ruta_nueva)
    ws_proc = wb_proc["CONFIGURACIÓN DE ANAQUEL"]
    
    print(f"  Buscando ELECTROLIFE ZERO POLVO NARANJA en fila 3 (después de borrar fila 1)...")
    encontrado = False
    for col in range(1, min(ws_proc.max_column + 1, 150)):
        val = ws_proc.cell(row=3, column=col).value
        if val and "NARANJA" in str(val).upper() and "POLVO" in str(val).upper():
            col_letter = get_column_letter(col)
            print(f"  ✓ ENCONTRADO en {col_letter} ({col}): {val}")
            print(f"    Fila 2 (SKU): {ws_proc.cell(row=2, column=col).value}")
            print(f"    Fila 4 (dato): {ws_proc.cell(row=4, column=col).value}")
            encontrado = True
            break
    
    if not encontrado:
        print(f"  ✗ NO ENCONTRADO - Mostrando todas las columnas con 'POLVO':")
        for col in range(1, min(ws_proc.max_column + 1, 150)):
            val = ws_proc.cell(row=3, column=col).value
            if val and "POLVO" in str(val).upper():
                col_letter = get_column_letter(col)
                print(f"    {col_letter}: {val}")
                print(f"      Fila 2: {ws_proc.cell(row=2, column=col).value}")
                print(f"      Fila 4: {ws_proc.cell(row=4, column=col).value}")
    
    # También buscar en PRODUCTOS
    print(f"\n[4] Verificando hoja PRODUCTOS...")
    ws_prod = wb_proc["PRODUCTOS"]
    print(f"  Buscando POLVO NARANJA en PRODUCTOS fila 4...")
    for col in range(1, min(ws_prod.max_column + 1, 150)):
        val = ws_prod.cell(row=4, column=col).value
        if val and "NARANJA" in str(val).upper() and "POLVO" in str(val).upper():
            col_letter = get_column_letter(col)
            print(f"  ✓ ENCONTRADO en {col_letter}: {val}")
            break
    
    print(f"\n[5] Archivo guardado en: {ruta_nueva}")
    
except Exception as e:
    print(f"  ERROR: {e}")
    import traceback
    traceback.print_exc()
