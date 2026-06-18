# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

print("="*90)
print("INVESTIGACIÓN: ELECTROLIFE ZERO POLVO NARANJA - ¿Dónde se pierde el dato?")
print("="*90)

# Cargar la MATRIZ ORIGINAL (respaldo - sin modificar)
ruta_original = r"C:\Users\katoH\Downloads\MATRIZ DE CATALOGACÓN actual 1 (4) (1) - copia.xlsx"
print(f"\n1. Abriendo MATRIZ ORIGINAL (respaldo): {ruta_original}")
wb_orig = load_workbook(ruta_original)
ws_orig = wb_orig["CONFIGURACIÓN DE ANAQUEL"]

# Buscar "ELECTROLIFE ZERO POLVO NARANJA" en la fila 4 (nombres de productos)
print("\nBuscando 'ELECTROLIFE ZERO POLVO NARANJA' en CONFIGURACIÓN DE ANAQUEL...")
col_naranja_original = None
for col in range(1, ws_orig.max_column + 1):
    val = ws_orig.cell(row=4, column=col).value
    if val and "NARANJA" in str(val).upper() and "POLVO" in str(val).upper():
        col_letter = get_column_letter(col)
        print(f"\n✓ ENCONTRADO EN COLUMNA {col_letter} ({col}): {val}")
        col_naranja_original = col
        
        # Mostrar datos de esa columna
        print(f"\n  Datos en {col_letter}:")
        print(f"    Fila 2 (Header): {ws_orig.cell(row=2, column=col).value}")
        print(f"    Fila 3 (SKU): {ws_orig.cell(row=3, column=col).value}")
        print(f"    Fila 4 (Nombre): {ws_orig.cell(row=4, column=col).value}")
        print(f"    Fila 5 (LUGAR 1): {ws_orig.cell(row=5, column=col).value}")
        print(f"    Fila 6 (LUGAR 2): {ws_orig.cell(row=6, column=col).value}")
        break

if not col_naranja_original:
    print("\n✗ NO ENCONTRADO")

# Ahora buscar en la matriz PROCESADA
print("\n" + "="*90)
print("2. Abriendo MATRIZ PROCESADA")
print("="*90)
ruta_procesada = r"C:\Users\katoH\Downloads\MATRIZ DE CATALOGACÓN actual 1 (4) (1).xlsx"
try:
    wb_proc = load_workbook(ruta_procesada)
    ws_proc = wb_proc["CONFIGURACIÓN DE ANAQUEL"]
    
    print("\nBuscando 'ELECTROLIFE ZERO POLVO NARANJA' en matriz procesada...")
    col_naranja_procesada = None
    for col in range(1, ws_proc.max_column + 1):
        val = ws_proc.cell(row=3, column=col).value  # Ahora fila 3 porque borramos fila 1
        if val and "NARANJA" in str(val).upper() and "POLVO" in str(val).upper():
            col_letter = get_column_letter(col)
            print(f"\n✓ ENCONTRADO EN COLUMNA {col_letter} ({col}): {val}")
            col_naranja_procesada = col
            
            print(f"\n  Datos en {col_letter} (después de borrar fila 1):")
            print(f"    Fila 1 (Header): {ws_proc.cell(row=1, column=col).value}")
            print(f"    Fila 2 (SKU): {ws_proc.cell(row=2, column=col).value}")
            print(f"    Fila 3 (Nombre): {ws_proc.cell(row=3, column=col).value}")
            print(f"    Fila 4 (LUGAR 1): {ws_proc.cell(row=4, column=col).value}")
            print(f"    Fila 5 (LUGAR 2): {ws_proc.cell(row=5, column=col).value}")
            break
    
    if not col_naranja_procesada:
        print("\n✗ NO ENCONTRADO en procesada")
        print("\n  Buscando parcialmente 'POLVO'...")
        for col in range(1, ws_proc.max_column + 1):
            val = ws_proc.cell(row=3, column=col).value
            if val and "POLVO" in str(val).upper():
                col_letter = get_column_letter(col)
                print(f"    {col_letter}: {val}")
                
except Exception as e:
    print(f"ERROR al abrir procesada: {e}")

# Verificar también la hoja PRODUCTOS
print("\n" + "="*90)
print("3. Verificando hoja PRODUCTOS (para ver columna vacía después de FIFTEEN PACK)")
print("="*90)

try:
    ws_prod = wb_proc["PRODUCTOS"]
    print(f"Hoja PRODUCTOS: {ws_prod.max_row} filas, {ws_prod.max_column} columnas")
    
    # Buscar FIFTEEN PACK
    print("\nBuscando 'ELECTROLIT FIFTEEN PACK MORA' en PRODUCTOS...")
    for col in range(1, ws_prod.max_column + 1):
        val = ws_prod.cell(row=4, column=col).value
        if val and "FIFTEEN" in str(val).upper() and "MORA" in str(val).upper():
            col_letter = get_column_letter(col)
            print(f"\n✓ ENCONTRADO FIFTEEN PACK EN {col_letter} ({col})")
            
            # Verificar columnas adyacentes
            print(f"\n  Contexto de columnas:")
            for c in range(max(1, col - 2), min(ws_prod.max_column + 1, col + 4)):
                col_ctx = get_column_letter(c)
                val_fila4 = ws_prod.cell(row=4, column=c).value
                val_fila5 = ws_prod.cell(row=5, column=c).value
                marker = " <-- FIFTEEN PACK" if c == col else ""
                marker_after = " <-- (VACÍA?)" if c == col + 1 else ""
                print(f"    {col_ctx}: F4={val_fila4}, F5={val_fila5}{marker}{marker_after}")
            break
    
except Exception as e:
    print(f"ERROR verificando PRODUCTOS: {e}")
