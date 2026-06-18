#!/usr/bin/env python3
import sys
sys.path.insert(0, r"c:\Users\katoH\OneDrive\Documentos\Proyectos personales\Automatizaciones storecheck L2")

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

salida = []

try:
    salida.append("Abriendo archivo...")
    ruta_archivo = r"C:\Users\katoH\Downloads\MATRIZ DE CATALOGACÓN actual 1 (5).xlsx"
    wb = load_workbook(ruta_archivo, data_only=False)
    ws = wb["CONFIGURACIÓN DE ANAQUEL"]
    
    salida.append(f"Hoja cargada. Max column: {ws.max_column}")
    
    # Buscar Fifteen y Zero
    fifteen_col = None
    zero_col = None
    
    salida.append("Buscando Fifteen y Zero...")
    for col in range(11, min(200, ws.max_column + 1)):
        valor = ws.cell(row=2, column=col).value
        if valor:
            valor_str = str(valor)
            if "Fifteen" in valor_str:
                fifteen_col = col
                salida.append(f"  Encontrado Fifteen en {get_column_letter(col)} ({col}): {valor_str[:40]}")
            elif "Zero" in valor_str and "Polvo" in valor_str:
                zero_col = col
                salida.append(f"  Encontrado Zero en {get_column_letter(col)} ({col}): {valor_str[:40]}")
    
    if fifteen_col and zero_col:
        salida.append(f"\nRango entre Fifteen ({get_column_letter(fifteen_col)}) y Zero ({get_column_letter(zero_col)}):")
        vacias = 0
        
        for col in range(fifteen_col, zero_col + 1):
            letra = get_column_letter(col)
            valor_f2 = ws.cell(row=2, column=col).value
            
            # Verificar filas 4, 5, 6
            tiene_datos_456 = False
            for fila in [4, 5, 6]:
                dato = ws.cell(row=fila, column=col).value
                if dato:
                    tiene_datos_456 = True
                    break
            
            if valor_f2:
                salida.append(f"  {letra:3s} ({col:3d}): MARCADOR: {str(valor_f2)[:35]}")
            elif tiene_datos_456:
                salida.append(f"  {letra:3s} ({col:3d}): vacío en fila 2, datos en 4-6")
            else:
                salida.append(f"  {letra:3s} ({col:3d}): *** COMPLETAMENTE VACÍO ***")
                vacias += 1
        
        salida.append(f"\nRESULTADO:")
        if vacias == 0:
            salida.append("SUCCESS: Cero columnas completamente vacías entre Fifteen y Zero")
        elif vacias == 1:
            salida.append(f"WARNING: Hay {vacias} columna vacía")
        else:
            salida.append(f"PROBLEM: Hay {vacias} columnas completamente vacías")
    else:
        salida.append("ERROR: No se encontraron Fifteen o Zero")
        if fifteen_col:
            salida.append(f"  Fifteen: {get_column_letter(fifteen_col)} ({fifteen_col})")
        if zero_col:
            salida.append(f"  Zero: {get_column_letter(zero_col)} ({zero_col})")

except Exception as e:
    salida.append(f"ERROR: {e}")
    import traceback
    salida.append(traceback.format_exc())

# Guardar resultado
with open(r"c:\Users\katoH\OneDrive\Documentos\Proyectos personales\Automatizaciones storecheck L2\resultado_verificacion.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(salida))

print("\n".join(salida))
