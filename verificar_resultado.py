#!/usr/bin/env python3
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import sys

try:
    # Abrir el archivo Excel que fue procesado
    ruta_archivo = r"C:\Users\katoH\Downloads\MATRIZ DE CATALOGACÓN actual 1 (5).xlsx"
    print("Abriendo archivo...")
    wb = load_workbook(ruta_archivo, data_only=False)
    ws = wb["CONFIGURACIÓN DE ANAQUEL"]
    print("Archivo abierto. Analizando...\n")

    # Buscar Fifteen y Zero
    columnas_interes = {}
    for col in range(11, min(ws.max_column + 1, 300)):
        valor = ws.cell(row=2, column=col).value
        if valor:
            valor_str = str(valor).lower()
            if "fifteen" in valor_str:
                columnas_interes['Fifteen'] = (col, get_column_letter(col), valor)
            elif "zero" in valor_str and "polvo" in valor_str:
                columnas_interes['Zero'] = (col, get_column_letter(col), valor)

    print("Columnas encontradas:")
    if 'Fifteen' in columnas_interes:
        col_idx, letra, valor = columnas_interes['Fifteen']
        print(f"  Fifteen: Col {letra} (idx {col_idx}): {valor}")
    if 'Zero' in columnas_interes:
        col_idx, letra, valor = columnas_interes['Zero']
        print(f"  Zero: Col {letra} (idx {col_idx}): {valor}")

    if 'Fifteen' in columnas_interes and 'Zero' in columnas_interes:
        col_fifteen = columnas_interes['Fifteen'][0]
        col_zero = columnas_interes['Zero'][0]
        
        print(f"\nColumnas entre Fifteen ({get_column_letter(col_fifteen)}) y Zero ({get_column_letter(col_zero)}):")
        for col in range(col_fifteen, col_zero + 1):
            letra = get_column_letter(col)
            valor = ws.cell(row=2, column=col).value
            datos_456 = [ws.cell(row=r, column=col).value for r in range(4, 7)]
            es_vacio_completo = not valor and not any(datos_456)
            
            if es_vacio_completo:
                print(f"  {letra:3s} (idx {col:3d}): [VACIO]")
            else:
                print(f"  {letra:3s} (idx {col:3d}): {str(valor or '<vacio>')[:30]}")
        
        # Contar cuantas columnas vacías hay entre Fifteen y Zero
        vacias = 0
        for col in range(col_fifteen + 1, col_zero):
            valor = ws.cell(row=2, column=col).value
            datos_456 = [ws.cell(row=r, column=col).value for r in range(4, 7)]
            if not valor and not any(datos_456):
                vacias += 1
        
        print(f"\nRESULTADO FINAL:")
        print(f"  Columnas completamente vacías entre Fifteen y Zero: {vacias}")
        if vacias == 0:
            print("  ✓ ¡ÉXITO! No hay columnas vacías entre los dos productos")
        elif vacias <= 2:
            print(f"  ⚠ Todavía hay {vacias} columnas vacías")
        else:
            print(f"  ✗ Hay {vacias} columnas vacías (problema no resuelto)")

except Exception as e:
    print(f"ERROR: {e}", file=sys.stderr)
    import traceback
    traceback.print_exc(file=sys.stderr)

