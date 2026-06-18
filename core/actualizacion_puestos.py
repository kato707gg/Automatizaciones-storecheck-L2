"""
Actualizacion de puestos.

Toma un libro de Excel que contenga las hojas Coordinador, Promotor y Supervisor,
filtra las filas que no correspondan a los perfiles permitidos en la columna 3 y
crea una hoja nueva llamada Cruces con la combinacion pedida.
"""

from __future__ import annotations

import os
import re
import unicodedata
from copy import copy

import openpyxl
from openpyxl.styles import Font, PatternFill


_ALLOWED_PROFILES = {
    "PROMOTOR COOR- PRO",
    "PROMOTOR- PRO",
    "SUPERVISOR- PRO",
    "COORDINADOR- PRO",
}

_REQUIRED_SHEETS = (
    "Coordinador",
    "Promotor",
    "Supervisor",
)

_SHEET_CONFIG = {
    "COORDINADOR": {
        "sheet": "Coordinador",
        "ruta_col": 14,
        "encargado_header": "Gerente",
    },
    "PROMOTOR": {
        "sheet": "Promotor",
        "ruta_col": 16,
        "encargado_header": "Supervisor",
    },
    "SUPERVISOR": {
        "sheet": "Supervisor",
        "ruta_col": 15,
        "encargado_header": "Coordinador",
    },
}


def _normalizar_texto(valor: object) -> str:
    texto = unicodedata.normalize("NFKD", str(valor))
    texto = "".join(ch for ch in texto if not unicodedata.combining(ch))
    texto = texto.replace("–", "-").replace("—", "-")
    texto = re.sub(r"\s*-\s*", "-", texto)
    texto = " ".join(texto.strip().split())
    return texto.upper()


def _normalizar_hoja(nombre: str) -> str:
    return _normalizar_texto(nombre)


def _normalizar_profesional(valor: object) -> str:
    return _normalizar_texto(valor)


def _perfil_permitido(valor: object) -> bool:
    if valor is None:
        return False
    normalizado = _normalizar_profesional(valor)
    for permitido in _ALLOWED_PROFILES:
        permitido_norm = _normalizar_texto(permitido)
        if normalizado == permitido_norm:
            return True
        if normalizado and permitido_norm.startswith(normalizado):
            return True
        if normalizado and normalizado.startswith(permitido_norm):
            return True
    return False


def _buscar_hoja(wb: openpyxl.Workbook, nombre_hoja: str):
    objetivo = _normalizar_hoja(nombre_hoja)
    for nombre in wb.sheetnames:
        if _normalizar_hoja(nombre) == objetivo:
            return wb[nombre]
    return None


def _buscar_columna(ws, fila_encabezados: int, encabezados) -> int | None:
    if isinstance(encabezados, str):
        encabezados = (encabezados,)
    objetivos = [_normalizar_texto(h) for h in encabezados]
    fila = list(ws.iter_rows(
        min_row=fila_encabezados,
        max_row=fila_encabezados,
        values_only=True,
    ))
    if not fila:
        return None
    fila = fila[0]

    for col_idx, valor in enumerate(fila, start=1):
        if valor is None:
            continue
        valor_norm = _normalizar_texto(valor)
        if valor_norm in objetivos:
            return col_idx

    for col_idx, valor in enumerate(fila, start=1):
        if valor is None:
            continue
        valor_norm = _normalizar_texto(valor)
        for objetivo in objetivos:
            if objetivo in valor_norm or valor_norm in objetivo:
                return col_idx
    return None


def _detectar_fila_encabezados(ws, encabezados: tuple[str, ...], max_filas: int = 25) -> int:
    objetivos = {_normalizar_texto(enc) for enc in encabezados}
    mejor_fila = 1
    mejores_hits = 0

    for fila_idx, fila in enumerate(
        ws.iter_rows(min_row=1, max_row=min(max_filas, ws.max_row), values_only=True),
        start=1,
    ):
        valores = {_normalizar_texto(valor) for valor in fila if valor is not None}
        hits = len(objetivos & valores)
        if hits > mejores_hits:
            mejores_hits = hits
            mejor_fila = fila_idx

    return mejor_fila


def _valor_col(fila, col_idx: int | None):
    if col_idx is None:
        return None
    if len(fila) < col_idx:
        return None
    return fila[col_idx - 1]


def _normalizar_correo(valor: object) -> str:
    return str(valor or "").strip().lower()


def _limpiar_duplicado(texto: object) -> str:
    if texto is None:
        return ""
    raw = str(texto).strip()
    if raw.upper().endswith(" DUPLICADO"):
        raw = raw[: -len(" DUPLICADO")].strip()
    return raw


def _subrayar_fila(ws, row_idx: int, col_inicio: int, col_fin: int):
    marcador = PatternFill(start_color="FFF59D", end_color="FFF59D", fill_type="solid")
    for col in range(col_inicio, col_fin + 1):
        celda = ws.cell(row=row_idx, column=col)
        celda.fill = marcador


def _crear_hoja_cruces(wb: openpyxl.Workbook):
    if "Cruces" in wb.sheetnames:
        wb.remove(wb["Cruces"])
    ws = wb.create_sheet("Cruces")
    encabezados = ["Puestos", "Correo", "Nombre completo", "Nombre", "Apellido", "Encargado"]
    for col_idx, valor in enumerate(encabezados, start=1):
        if valor is None:
            continue
        celda = ws.cell(row=1, column=col_idx, value=valor)
        celda.font = Font(bold=True)
    ws.freeze_panes = "A2"
    return ws


def _copiar_estilo_fila(origen, destino):
    for src_cell, dst_cell in zip(origen, destino):
        if src_cell.has_style:
            dst_cell._style = copy(src_cell._style)
        if src_cell.number_format:
            dst_cell.number_format = src_cell.number_format


def procesar_actualizacion_puestos(
    ruta_catalogo_puestos: str,
    ruta_layout_roles: str,
    carpeta_salida: str,
) -> str:
    """Procesa el libro y devuelve la ruta del archivo generado."""
    if not ruta_catalogo_puestos or not os.path.exists(ruta_catalogo_puestos):
        raise ValueError("No se encontro el catalogo de puestos.")
    if not ruta_layout_roles or not os.path.exists(ruta_layout_roles):
        raise ValueError("No se encontro la plantilla layout_roles.")

    wb_data = None
    wb_out = None

    try:
        wb_data = openpyxl.load_workbook(
            ruta_catalogo_puestos,
            data_only=True,
            read_only=True,
        )
        wb_out = openpyxl.load_workbook(ruta_layout_roles, keep_links=False)

        hojas = {
            key: _buscar_hoja(wb_data, config["sheet"])
            for key, config in _SHEET_CONFIG.items()
        }

        for key, ws in hojas.items():
            if ws is None:
                disponibles = ", ".join(wb_data.sheetnames)
                raise ValueError(
                    f"No se encontro la hoja requerida: {_SHEET_CONFIG[key]['sheet']}. "
                    f"Hojas disponibles en catalogo: {disponibles}"
                )

        ws_cruces = _crear_hoja_cruces(wb_out)

        filas_cruces = []

        for key in ("COORDINADOR", "PROMOTOR", "SUPERVISOR"):
            config = _SHEET_CONFIG[key]
            ws = hojas[key]

            fila_enc = _detectar_fila_encabezados(
                ws,
                (
                    "Ruta",
                    "RUTA",
                    "Correo electrónico",
                    "Correo electronico",
                    "Correo",
                    "Nombre completo",
                    "Nombre",
                    "Apellido",
                    "Puesto",
                    "Perfil",
                    "Nomina",
                    "Nómina",
                    config["encargado_header"],
                ),
                max_filas=60,
            )
            fila_datos = fila_enc + 1

            col_ruta = _buscar_columna(ws, fila_enc, ("Ruta", "RUTA")) or config["ruta_col"]
            col_correo = _buscar_columna(ws, fila_enc, (
                "Correo electrónico", "Correo electronico", "Correo"
            )) or 5
            col_nombre_completo = _buscar_columna(ws, fila_enc, ("Nombre completo", "Nombre Completo"))
            col_nombre = _buscar_columna(ws, fila_enc, ("Nombre", "Nombre(s)")) or 8
            col_apellido = _buscar_columna(ws, fila_enc, ("Apellido", "Apellidos")) or 9
            col_perfil = _buscar_columna(ws, fila_enc, ("Perfil",))
            col_puesto = _buscar_columna(ws, fila_enc, ("Puesto",))
            col_encargado = _buscar_columna(ws, fila_enc, config["encargado_header"])

            for fila in ws.iter_rows(min_row=fila_datos, values_only=True):
                perfil = _valor_col(fila, col_perfil)
                puesto = _valor_col(fila, col_puesto)
                perfil_fallback = fila[2] if len(fila) >= 3 else None
                if not (_perfil_permitido(perfil) or _perfil_permitido(puesto) or _perfil_permitido(perfil_fallback)):
                    continue

                puestos = _valor_col(fila, col_ruta)
                correo = _valor_col(fila, col_correo)
                nombre_completo = _valor_col(fila, col_nombre_completo)
                nombre = _valor_col(fila, col_nombre)
                apellido = _valor_col(fila, col_apellido)
                encargado = _valor_col(fila, col_encargado)

                if not nombre_completo:
                    partes = [p for p in (nombre, apellido) if p not in (None, "")]
                    if partes:
                        nombre_completo = " ".join(str(p).strip() for p in partes if str(p).strip())

                filas_cruces.append([
                    puestos,
                    correo,
                    nombre_completo,
                    nombre,
                    apellido,
                    encargado,
                ])

        if not filas_cruces:
            raise ValueError(
                "No se encontraron filas para Cruces. "
                "Revisa los encabezados y los valores de Perfil/Puesto en el catalogo."
            )

        for row_idx, valores in enumerate(filas_cruces, start=2):
            for col_idx, valor in enumerate(valores, start=1):
                if valor is None:
                    continue
                ws_cruces.cell(row=row_idx, column=col_idx, value=valor)

        # ── Integracion con hoja Puestos en plantilla ──────────────
        if "Puestos" not in wb_out.sheetnames:
            raise ValueError("No se encontro la hoja 'Puestos' en la plantilla layout_roles.")

        ws_puestos = wb_out["Puestos"]
        fila_enc_p = _detectar_fila_encabezados(
            ws_puestos,
            ("Puesto", "Correo electrónico", "Correo electronico", "Correo"),
            max_filas=60,
        )
        fila_datos_p = fila_enc_p + 1

        col_puesto_p = _buscar_columna(ws_puestos, fila_enc_p, "Puesto") or 2
        col_correo_p = _buscar_columna(ws_puestos, fila_enc_p, (
            "Correo electrónico", "Correo electronico", "Correo"
        )) or 5
        col_edit = 22

        # Mapa correo -> puesto desde Cruces
        mapa_cruces = {}
        for fila in ws_cruces.iter_rows(min_row=2, values_only=True):
            correo = _normalizar_correo(fila[1]) if len(fila) > 1 else ""
            puesto = fila[0] if len(fila) > 0 else None
            if correo:
                mapa_cruces[correo] = puesto

        puestos_catalogo = set()
        for fila in ws_puestos.iter_rows(min_row=fila_datos_p, values_only=True):
            puesto_actual = _limpiar_duplicado(_valor_col(fila, col_puesto_p))
            if puesto_actual:
                puestos_catalogo.add(_normalizar_texto(puesto_actual))

        max_row_p = ws_puestos.max_row
        for row_idx in range(fila_datos_p, max_row_p + 1):
            correo_val = ws_puestos.cell(row=row_idx, column=col_correo_p).value
            if correo_val in (None, ""):
                continue

            correo_norm = _normalizar_correo(correo_val)
            puesto_en_cruces = mapa_cruces.get(correo_norm)

            if puesto_en_cruces in (None, ""):
                continue

            puesto_actual_raw = ws_puestos.cell(row=row_idx, column=col_puesto_p).value
            puesto_actual_base = _limpiar_duplicado(puesto_actual_raw)
            if puesto_actual_base != str(puesto_en_cruces).strip():
                ws_puestos.cell(row=row_idx, column=col_puesto_p, value=puesto_en_cruces)
                ws_puestos.cell(row=row_idx, column=col_edit, value="EDIT")

        # ── Subrayar puestos nuevos en Cruces ───────────────────────
        for row_idx in range(2, ws_cruces.max_row + 1):
            puesto_cruces = ws_cruces.cell(row=row_idx, column=1).value
            if not puesto_cruces:
                continue
            if _normalizar_texto(puesto_cruces) not in puestos_catalogo:
                _subrayar_fila(ws_cruces, row_idx, 1, 6)

        # Ajuste simple de anchos para facilitar la lectura.
        anchuras = {
            "A": 24,
            "B": 32,
            "C": 30,
            "D": 24,
            "E": 24,
            "F": 24,
        }
        for columna, ancho in anchuras.items():
            ws_cruces.column_dimensions[columna].width = ancho

        os.makedirs(carpeta_salida, exist_ok=True)
        base = os.path.splitext(os.path.basename(ruta_layout_roles))[0]
        ext = os.path.splitext(ruta_layout_roles)[1] or ".xlsx"
        nombre_salida = f"{base}_actualizado{ext}"
        ruta_salida = os.path.join(carpeta_salida, nombre_salida)

        wb_out.save(ruta_salida)
        return ruta_salida
    finally:
        if wb_data is not None:
            wb_data.close()
        if wb_out is not None:
            wb_out.close()
