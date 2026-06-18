from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

path = r'C:\Users\katoH\Downloads\Matriz_Aqualit_CF_20260610_044753.xlsx'
wb = load_workbook(path, data_only=False)
needle = 'NARANJA'

for sheet_name in ['CONFIGURACIÓN DE ANAQUEL', 'PRODUCTOS']:
    ws = wb[sheet_name]
    print(sheet_name)
    matches = []
    for row in range(1, min(ws.max_row + 1, 15)):
        for col in range(1, min(ws.max_column + 1, 220)):
            value = ws.cell(row=row, column=col).value
            if value and needle in str(value).upper():
                matches.append((get_column_letter(col), row, value, ws.cell(row=2, column=col).value))
    if matches:
        for letter, row, value, sku in matches[:20]:
            print(letter, row, sku, value)
    else:
        print('NOT FOUND')
